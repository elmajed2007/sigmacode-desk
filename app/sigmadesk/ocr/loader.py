"""Turn a file into pages of text: PDF text layer when usable, OCR otherwise.

Supported: PDF, images (png jpg jpeg tif tiff bmp webp), DOCX, TXT/MD.
Each page: {"page", "text", "method": "text"|"ocr", "confidence", "lines": [...]}
"""
from __future__ import annotations

import re
import unicodedata
from pathlib import Path

import numpy as np

from . import arabic

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
TEXT_EXT = {".txt", ".md", ".csv"}
SUPPORTED = IMAGE_EXT | TEXT_EXT | {".pdf", ".docx"}

_WORDISH = re.compile(r"[\w؀-ۿ]{2,}")


class LoadError(Exception):
    pass


def kind_of(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return "pdf"
    if ext in IMAGE_EXT:
        return "image"
    if ext == ".docx":
        return "docx"
    if ext in TEXT_EXT:
        return "text"
    raise LoadError(f"Unsupported file type: {ext or filename}")


def clean_text_layer(text: str) -> str:
    """PDF text layers often carry presentation forms and odd spacing; Arabic lines
    are sometimes stored in visual order - fixed per page by the «ال» test."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [arabic.normalize_arabic(unicodedata.normalize("NFC", l)).strip() for l in text.split("\n")]
    lines = [re.sub(r"[ \t ]+", " ", l) for l in lines if l.strip()]
    if lines and arabic.needs_reverse(lines):
        lines = [arabic.pred_reverse(l) if arabic.is_rtl(l) else l for l in lines]
    return "\n".join(lines)


def usable_text(text: str) -> bool:
    """Enough real words (not a scanned page with a few stray glyphs)."""
    if not text:
        return False
    words = _WORDISH.findall(text)
    bad = sum(1 for c in text if c == "�" or 0xE000 <= ord(c) <= 0xF8FF)
    return len(words) >= 8 and bad <= 0.02 * len(text)


def pdf_page_count(path: Path) -> int:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(path))
    try:
        return len(pdf)
    finally:
        pdf.close()


def render_pdf_page(path: Path, index: int, scale: float = 2.0) -> np.ndarray:
    """BGR image of a PDF page (pdfium renders BGRx)."""
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(path))
    try:
        return pdf[index].render(scale=scale).to_numpy()[:, :, :3].copy()
    finally:
        pdf.close()


def page_png(path: Path, kind: str, page_no: int, scale: float = 1.5) -> bytes:
    """PNG bytes of one page for the viewer."""
    import cv2

    if kind == "pdf":
        img = render_pdf_page(path, page_no - 1, scale)
    elif kind == "image":
        img = load_images(path)[page_no - 1]
    else:
        raise LoadError("This document has no page images.")
    ok, buf = cv2.imencode(".png", img)
    if not ok:
        raise LoadError("Could not render the page.")
    return buf.tobytes()


def load_images(path: Path) -> list[np.ndarray]:
    """All frames of an image file as BGR arrays (multi-page TIFF supported)."""
    from PIL import Image, ImageSequence

    out = []
    with Image.open(path) as im:
        for frame in ImageSequence.Iterator(im):
            rgb = np.array(frame.convert("RGB"))
            out.append(rgb[:, :, ::-1].copy())
    return out


def docx_text(path: Path) -> str:
    import docx

    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1256", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def split_pages(text: str, size: int = 3000) -> list[str]:
    """Long plain texts become pseudo-pages (on paragraph breaks) for citations."""
    paras, pages, cur = text.split("\n"), [], []
    n = 0
    for p in paras:
        if n + len(p) > size and cur:
            pages.append("\n".join(cur))
            cur, n = [], 0
        cur.append(p)
        n += len(p) + 1
    if cur:
        pages.append("\n".join(cur))
    return [p for p in pages if p.strip()] or [""]


def extract(path: Path, kind: str, engine, ocr_settings, *, pdf_scale=2.0, force_ocr=False,
            max_pages=500, progress=None) -> list[dict]:
    """Pages of the document. ``engine``: OcrEngine (loaded lazily by the caller)."""
    path = Path(path)
    if kind == "text":
        return [{"page": i, "text": t, "method": "text", "confidence": None, "lines": []}
                for i, t in enumerate(split_pages(read_text_file(path)), start=1)]
    if kind == "docx":
        return [{"page": i, "text": t, "method": "text", "confidence": None, "lines": []}
                for i, t in enumerate(split_pages(docx_text(path)), start=1)]
    if kind == "image":
        images = load_images(path)
        if len(images) > max_pages:
            raise LoadError(f"The file has {len(images)} pages; the limit is {max_pages}.")
        return _ocr(engine(), images, list(range(1, len(images) + 1)), ocr_settings, progress)
    if kind != "pdf":
        raise LoadError("Unsupported document kind.")

    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(path))
    try:
        n = len(pdf)
        if n > max_pages:
            raise LoadError(f"The PDF has {n} pages; the limit is {max_pages}.")
        pages, to_ocr = {}, []
        for i in range(n):
            text = ""
            if not force_ocr:
                tp = pdf[i].get_textpage()
                try:
                    text = clean_text_layer(tp.get_text_range())
                finally:
                    tp.close()
            if usable_text(text):
                pages[i + 1] = {"page": i + 1, "text": text, "method": "text", "confidence": None, "lines": []}
            else:
                to_ocr.append(i)
    finally:
        pdf.close()
    if to_ocr:
        images = [render_pdf_page(path, i, pdf_scale) for i in to_ocr]
        for p in _ocr(engine(), images, [i + 1 for i in to_ocr], ocr_settings, progress):
            pages[p["page"]] = p
    elif progress:
        progress(1, 1)
    return [pages[k] for k in sorted(pages)]


def _ocr(engine, images, page_numbers, settings, progress):
    raw = []
    for k, (img, no) in enumerate(zip(images, page_numbers), start=1):
        raw.append(engine.ocr_image(img, settings, no))
        if progress:
            progress(k, len(images))
    built = arabic.build_pages(raw, settings.reverse_mode)
    return [{"page": p["page"], "text": p["text"], "method": "ocr",
             "confidence": (sum(l["score"] * len(l["text"]) for l in p["lines"])
                            / max(1, sum(len(l["text"]) for l in p["lines"]))) if p["lines"] else 0.0,
             "lines": [{"text": l["text"], "score": l["score"], "bbox": l["bbox"]} for l in p["lines"]]}
            for p in built["pages"]]
